"""Weekly PowerPoint automation for FastPark.

Behavior:
- Refreshes the slide data workbook tabs from latest forecast accuracy output.
- Archives the current deck with a date stamp.
"""

from __future__ import annotations

import contextlib
from datetime import datetime
from pathlib import Path
import shutil

from openpyxl import load_workbook


def to_date_key(value):
    """Normalize datetime-like values to date keys for robust joins."""
    if hasattr(value, "date"):
        return value.date()
    return value


def find_label_row(ws, label: str, start_col: int) -> int:
    """Find the row where a section label appears in a specific column."""
    for row in range(1, ws.max_row + 1):
        cell_val = ws.cell(row=row, column=start_col).value
        if isinstance(cell_val, str) and cell_val.strip().lower() == label.lower():
            return row
    raise ValueError(f"Label '{label}' not found in sheet '{ws.title}' (col {start_col}).")


def read_block(ws, label: str, start_col: int) -> tuple[list, dict]:
    """Read a labelled table block and return ordered dates + value map."""
    label_row = find_label_row(ws, label, start_col)
    dates = []
    values_map = {}

    row = label_row + 1
    while row <= ws.max_row:
        date_val = ws.cell(row=row, column=start_col).value
        if date_val is None:
            break

        horizon_vals = [
            ws.cell(row=row, column=start_col + offset).value
            for offset in range(1, 6)
        ]

        dates.append(date_val)
        values_map[to_date_key(date_val)] = horizon_vals
        row += 1

    return dates, values_map


def choose_slide_data_workbook(slides_dir: Path) -> Path:
    """Pick the workbook that feeds slide charts."""
    preferred = slides_dir / "FastPark Weekly Side Data.xlsx"
    if preferred.exists():
        return preferred

    candidates = [
        p for p in slides_dir.glob("*.xlsx")
        if p.is_file() and not p.name.startswith("~$")
    ]

    if not candidates:
        raise FileNotFoundError(
            f"No .xlsx workbook found in slides folder: {slides_dir}"
        )

    # Fall back to most recently modified workbook.
    return max(candidates, key=lambda p: p.stat().st_mtime)


def choose_latest_forecast_accuracy_file(accuracy_dir: Path) -> Path:
    """Select the latest forecast accuracy workbook output."""
    candidates = sorted(
        accuracy_dir.glob("forecast_accuracy_*.xlsx"),
        key=lambda p: p.stat().st_mtime,
    )
    if not candidates:
        raise FileNotFoundError(
            f"No forecast accuracy files found in: {accuracy_dir}"
        )
    return candidates[-1]


def clear_data_rows(ws, first_data_row: int, last_data_row: int, max_col: int) -> None:
    """Clear existing table data rows but leave headers/format scaffolding."""
    for row in range(first_data_row, last_data_row + 1):
        for col in range(1, max_col + 1):
            ws.cell(row=row, column=col).value = None


def write_accuracy_tab(target_ws, rows: list[tuple]) -> None:
    """Write rows to target tab with schema: Date, Actual, 21,14,7,3,1."""
    first_data_row = 2
    max_col = 7
    last_existing_row = max(target_ws.max_row, first_data_row + len(rows) + 5)

    clear_data_rows(target_ws, first_data_row, last_existing_row, max_col)

    for idx, row_vals in enumerate(rows, start=first_data_row):
        for col_idx, value in enumerate(row_vals, start=1):
            target_ws.cell(row=idx, column=col_idx).value = value


def refresh_slide_data_workbook(source_accuracy_file: Path, slide_data_file: Path) -> None:
    """Refresh Entries/Exits tabs from latest forecast accuracy output."""
    source_wb = load_workbook(source_accuracy_file, data_only=False)
    if "Forecast Accuracy" not in source_wb.sheetnames:
        raise ValueError(
            f"'Forecast Accuracy' sheet not found in {source_accuracy_file}"
        )

    src = source_wb["Forecast Accuracy"]

    entry_dates, entry_forecast = read_block(src, "Entries Forecast", 1)
    _, entry_actual = read_block(src, "Entries Actuals", 1)

    exit_dates, exit_forecast = read_block(src, "Exits Forecast", 10)
    _, exit_actual = read_block(src, "Exits Actuals", 10)

    target_wb = load_workbook(slide_data_file, data_only=False)

    if "Entries Accuracy" not in target_wb.sheetnames:
        raise ValueError(f"'Entries Accuracy' sheet not found in {slide_data_file}")
    if "Exits Accuracy" not in target_wb.sheetnames:
        raise ValueError(f"'Exits Accuracy' sheet not found in {slide_data_file}")

    entry_rows = []
    for date_val in entry_dates:
        key = to_date_key(date_val)
        entry_rows.append((
            date_val,
            entry_actual.get(key, [None] * 5)[0],
            *entry_forecast.get(key, [None] * 5),
        ))

    exit_rows = []
    for date_val in exit_dates:
        key = to_date_key(date_val)
        exit_rows.append((
            date_val,
            exit_actual.get(key, [None] * 5)[0],
            *exit_forecast.get(key, [None] * 5),
        ))

    write_accuracy_tab(target_wb["Entries Accuracy"], entry_rows)
    write_accuracy_tab(target_wb["Exits Accuracy"], exit_rows)

    target_wb.save(slide_data_file)

    print(f"Updated slide data workbook: {slide_data_file}")
    print(f"Entries rows written: {len(entry_rows)}")
    print(f"Exits rows written: {len(exit_rows)}")


def choose_deck(slides_dir: Path) -> Path:
    """Return the most recently modified deck in the slides directory."""
    candidates = [
        p for p in slides_dir.glob("*.pptx")
        if p.is_file() and not p.name.startswith("~$")
    ]

    if not candidates:
        raise FileNotFoundError(
            f"No .pptx deck found in slides folder: {slides_dir}"
        )

    # Use the latest modified deck to support template version updates.
    return max(candidates, key=lambda p: p.stat().st_mtime)


def build_archive_path(archive_dir: Path, source_deck: Path) -> Path:
    """Create a dated archive filename (one archive file per day)."""
    date_stamp = datetime.now().strftime("%Y-%m-%d")
    return archive_dir / f"{source_deck.stem}_{date_stamp}{source_deck.suffix}"


def refresh_powerpoint_links(deck_path: Path) -> None:
    """Refresh linked content in a PowerPoint file via COM automation."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        print("pywin32 is not installed. Skipping PowerPoint link refresh.")
        return

    app = None
    pres = None
    refreshed_shapes = 0
    refreshed_by_slide: dict[int, int] = {}

    try:
        app = win32com.client.Dispatch("PowerPoint.Application")
        app.Visible = 1

        pres = app.Presentations.Open(
            str(deck_path),
            ReadOnly=False,
            Untitled=False,
            WithWindow=False,
        )

        # Refresh global links first.
        with contextlib.suppress(Exception):
            pres.UpdateLinks()

        # Refresh shape-level linked objects and charts.
        for slide in pres.Slides:
            slide_num = int(slide.SlideIndex)
            refreshed_by_slide.setdefault(slide_num, 0)

            for shape in slide.Shapes:
                updated = False

                with contextlib.suppress(Exception):
                    shape.LinkFormat.Update()
                    updated = True

                if not updated:
                    with contextlib.suppress(Exception):
                        if shape.HasChart:
                            with contextlib.suppress(Exception):
                                shape.Chart.Refresh()
                            updated = True

                if updated:
                    refreshed_shapes += 1
                    refreshed_by_slide[slide_num] += 1

        pres.Save()
        print(f"PowerPoint links refreshed: {deck_path}")
        print(f"Linked/chart shapes refreshed: {refreshed_shapes}")
        per_slide = ", ".join(
            f"S{slide}:{count}"
            for slide, count in sorted(refreshed_by_slide.items())
            if count > 0
        )
        if per_slide:
            print(f"Per-slide refresh counts: {per_slide}")
        else:
            print("Per-slide refresh counts: none")

    except Exception as exc:
        print(f"PowerPoint auto-refresh warning: {exc}")

    finally:
        if pres is not None:
            with contextlib.suppress(Exception):
                pres.Close()

        if app is not None:
            with contextlib.suppress(Exception):
                app.Quit()


def main() -> int:
    comparison_root = Path(__file__).resolve().parent
    short_term_models_root = comparison_root.parents[1]

    outputs_root = short_term_models_root / "Outputs"
    slides_dir = outputs_root / "Weekly Forecast Slides"
    archive_dir = slides_dir / "Archive"
    accuracy_dir = comparison_root / "Forecast accuracy"

    if not slides_dir.exists():
        raise FileNotFoundError(
            f"Slides folder not found: {slides_dir}"
        )

    archive_dir.mkdir(parents=True, exist_ok=True)

    source_accuracy_file = choose_latest_forecast_accuracy_file(accuracy_dir)
    slide_data_file = choose_slide_data_workbook(slides_dir)
    refresh_slide_data_workbook(source_accuracy_file, slide_data_file)

    source_deck = choose_deck(slides_dir)
    refresh_powerpoint_links(source_deck)
    archive_path = build_archive_path(archive_dir, source_deck)

    shutil.copy2(source_deck, archive_path)

    print("Weekly PowerPoint step complete")
    print(f"Source accuracy: {source_accuracy_file}")
    print(f"Slide data file: {slide_data_file}")
    print(f"Source deck:  {source_deck}")
    print(f"Archive copy: {archive_path}")
    print("Note: this updates the Excel data source and archives the deck.")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
